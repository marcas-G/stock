"""财报 xlsx 快照解析（Plan P T8）：`*更新简化个股基本面数据.xlsx` → polars 帧 + fact parquet。

源（夸克网盘 `财报报表---有史以来--每周更新` → `data/raw/financial/`）：sheet1、51 列，
实测样本列头见 `lib/fundamentals.py::SRC_TO_OUT`（逐字映射；`code/更新日期/…/未分配利润`）。
输出 30 列（列序=`lib.fundamentals.OUT_COLUMNS`），单位沿用源（股本=万股、金额=元），
日期 `YYYYMMDD`→`datetime.date`。

**限制声明（T8 裁决）**：输出是**当期快照**（每行一个源更新日），非历史 PIT 序列；
PIT 历史待多期快照累积或人工 `*_financial.parquet`（manual_required）。
`updated_date` 缺失的行（源侧占位/停更行，实测 16 行）不属任何快照，解析时丢弃并计数；
否则 CH 主键 `(updated_date, ts_code)` 无法成立（CH 排序键不允许 Nullable）。

fact 落盘：`data/fact/fundamentals/fundamentals_snapshot.parquet`，覆盖写；旧版留 1 份
`*.prev`（tmp+os.replace 原子写）。CH 灌入由 `ch_ingest/ingest_fundamentals.py` 消费 fact。

用法：
    from pan_update import parse_fundamentals_xlsx as fp
    df = fp.parse_xlsx("data/raw/financial/2026-09-04更新简化个股基本面数据.xlsx")
    fp.write_fact(df, fp.DEFAULT_FACT)
"""
from __future__ import annotations

import argparse
import datetime
import os
import sys
from pathlib import Path

import openpyxl
import polars as pl

try:
    from lib import fundamentals as FU
except ModuleNotFoundError:  # 脚本直启（无 conftest 铺路）→ 补 platform/tools 再导入
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from lib import fundamentals as FU

from pan_update import config

OUT_COLUMNS = FU.OUT_COLUMNS
NUM_COLUMNS = FU.NUM_COLUMNS
OUT_SCHEMA = FU.OUT_SCHEMA

DEFAULT_SRC = config.repo_root() / "data" / "raw" / "financial"
DEFAULT_FACT = config.repo_root() / "data" / "fact" / "fundamentals" / \
    "fundamentals_snapshot.parquet"

_BLANK = {"", "None", "-", "—"}


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _to_date(v) -> datetime.date | None:
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    t = str(v).strip()
    if t in _BLANK:
        return None
    if len(t) == 8 and t.isdigit():
        return datetime.date(int(t[:4]), int(t[4:6]), int(t[6:8]))
    raise ValueError(f"无法解析日期: {v!r}")


def _to_float(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip()
    if t in _BLANK:
        return None
    try:
        return float(t)
    except ValueError as ex:
        raise ValueError(f"无法解析数值: {v!r}") from ex


def _to_str(v) -> str | None:
    if v is None:
        return None
    t = str(v).strip()
    return None if t in _BLANK else t


def parse_rows(path: Path) -> tuple[list[dict], int]:
    """xlsx → (行字典列表, 丢弃的 updated_date 缺失行数)；表头缺列/值非法 loud fail。"""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"财报 xlsx 不存在：{path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        raise ValueError("财报表头缺失：空工作簿")
    header = ["" if h is None else str(h).strip() for h in header]
    missing = [src for src, _ in FU.SRC_TO_OUT if src not in header]
    if missing:
        raise ValueError(f"财报表头缺列：{missing}")
    pos = {src: header.index(src) for src, _ in FU.SRC_TO_OUT}

    rows: list[dict] = []
    dropped = 0
    for raw in it:
        if raw is None or not any(not _blank(c) for c in raw):
            continue
        updated = _to_date(raw[pos["更新日期"]] if pos["更新日期"] < len(raw) else None)
        if updated is None:
            dropped += 1
            continue
        code6 = FU.code6_of(raw[pos["code"]] if pos["code"] < len(raw) else None)
        market = _to_str(raw[pos["市场"]] if pos["市场"] < len(raw) else None)
        out = {"ts_code": FU.ts_code_of(code6, market), "updated_date": updated}
        for src, name in FU.SRC_TO_OUT:
            if name in out:
                continue
            v = raw[pos[src]] if pos[src] < len(raw) else None
            if name in FU.DATE_COLUMNS:
                out[name] = _to_date(v)
            elif name in FU.TEXT_COLUMNS:
                out[name] = _to_str(v)
            else:
                out[name] = _to_float(v)
        rows.append(out)
    return rows, dropped


def parse_xlsx(path: Path) -> pl.DataFrame:
    """xlsx → 快照帧（30 列，schema=`lib.fundamentals.OUT_SCHEMA`；空表保留 schema）。"""
    rows, _ = parse_rows(path)
    return pl.DataFrame(rows, schema=FU.OUT_SCHEMA)


def write_fact(df: pl.DataFrame, out: Path) -> Path:
    """覆盖写 fact；已存在则先轮换为 `<name>.prev`（只留 1 份）；tmp+replace 原子。"""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        os.replace(out, out.with_name(out.name + ".prev"))
    tmp = out.with_name(out.name + ".tmp")
    df.write_parquet(tmp)
    os.replace(tmp, out)
    return out


def find_latest_xlsx(src: Path) -> Path:
    """单文件直通；目录取文件名最大者（源命名以 `YYYY-MM-DD` 前缀，字典序=时间序）。"""
    src = Path(src)
    if src.is_file():
        return src
    if not src.is_dir():
        raise FileNotFoundError(f"财报源不存在：{src}")
    cands = sorted(src.glob("*.xlsx"))
    if not cands:
        raise FileNotFoundError(f"财报源目录无 xlsx：{src}")
    return cands[-1]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--src", type=Path, default=None,
                    help=f"xlsx 文件或目录（缺省 {DEFAULT_SRC}）")
    ap.add_argument("--out", type=Path, default=None,
                    help=f"fact parquet 路径（缺省 {DEFAULT_FACT}）")
    args = ap.parse_args(argv)
    src = find_latest_xlsx(args.src or DEFAULT_SRC)
    rows, dropped = parse_rows(src)
    df = pl.DataFrame(rows, schema=FU.OUT_SCHEMA)
    out = write_fact(df, args.out or DEFAULT_FACT)
    print(f"  财报快照：{src.name} → {df.height:,} rows"
          f"（丢弃 updated_date 缺失 {dropped} 行）", flush=True)
    print(f"  fact: {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
